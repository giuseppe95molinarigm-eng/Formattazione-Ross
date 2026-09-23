#!/usr/bin/env python3
"""Automated QC for the Declutter Beyond interior PDF.

1. text integrity: every manuscript word appears, in order, in the PDF
2. no Toolkit/Checklist content leaked into the book
3. TOC page numbers match the pages where each opening actually falls
4. running heads / folios present only where expected
5. no heading stranded at the foot of a page
6. all fonts embedded, trim size 6x9
"""
import difflib
import re
import sys
from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).resolve().parent
PDF = HERE / "output" / "Declutter_Beyond_Interior_6x9.pdf"
MS = HERE / "source" / "Declutter_Beyond_Manuscript_v13_24_1.1.docx"
CHECKLIST = Path(sys.argv[1]) if len(sys.argv) > 1 else None

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
report = []


def say(s=""):
    print(s)
    report.append(s)


def words(s):
    s = s.replace("­", "").replace("​", "")
    s = re.sub(r"(\w)[-\u2010](\w)", r"\1\2", s)
    s = s.lower().replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    return re.findall(r"[a-z0-9$%'\"]+|[^\sa-z0-9]", s)


doc = pymupdf.open(PDF)
N = len(doc)

# ---------------------------------------------------------------- page text
def body_lines(page):
    """Text lines excluding running head (y<60) and foot folio (y>600).
    A drop cap is re-attached to the first line of its paragraph."""
    out, caps = [], []
    for b in page.get_text("dict")["blocks"]:
        if b["type"] != 0:
            continue
        for l in b["lines"]:
            y = l["spans"][0]["origin"][1]
            if y < 60 or y > 600:
                continue
            spans = [s for s in l["spans"] if s["text"].strip() or s["text"] == " "]
            if spans and spans[0]["size"] > 30 and "Garamond" in spans[0]["font"]:
                caps.append((y, spans[0]["text"].strip()))
                spans = spans[1:]
            t = "".join(s["text"] for s in spans)
            if t.strip():
                out.append([y, t, spans])
    out.sort(key=lambda x: x[0])
    for cy, letter in caps:
        cand = [ln for ln in out if cy - 24 < ln[0] < cy - 4]
        if cand:
            cand[0][1] = letter + cand[0][1]
    return [tuple(x) for x in out]


# find TOC pages (roman folios) and skip them for the integrity check
toc_pages = [i for i in range(N) if "TABLE OF CONTENTS" in doc[i].get_text() or
             (i in (3, 4) and "...." in doc[i].get_text())]

pdf_text = []
for i in range(N):
    if i in toc_pages:
        continue
    for y, t, spans in body_lines(doc[i]):
        pdf_text.append(t)
joined = "\n".join(pdf_text)
# undo end-of-line hyphenation introduced by justification
joined = re.sub(r"(\w)[-‐]\n(\w)", r"\1\2", joined)
pdf_words = words(joined)

d = docx.Document(str(MS))
ms_paras = [p for p in d.element.body if p.tag == W + "p"]
ms_text = []
for p in ms_paras:
    t = "".join(x.text or "" for x in p.iter(W + "t"))
    t = re.sub(r"^(CHAPTER \d+|Introduction|CONCLUSION|Appendix [AB]|Phase [IVX]+):\s*", r"\1 ", t)
    ms_text.append(t)
ms_words = words("\n".join(ms_text))

say("== 1. Text integrity (manuscript vs PDF, word sequence) ==")
say(f"manuscript words/tokens: {len(ms_words)}   PDF body tokens: {len(pdf_words)}")
sm = difflib.SequenceMatcher(None, ms_words, pdf_words, autojunk=False)
problems = 0
for tag, a1, a2, b1, b2 in sm.get_opcodes():
    if tag == "equal":
        continue
    ms_frag = " ".join(ms_words[a1:a2])
    pdf_frag = " ".join(pdf_words[b1:b2])
    problems += 1
    say(f"  {tag}: MS[{ms_frag[:120]}]  PDF[{pdf_frag[:120]}]")
say(f"difference blocks: {problems}")

# ---------------------------------------------------------------- toolkit leak
say("\n== 2. Toolkit / Checklist content not inserted ==")
if CHECKLIST and CHECKLIST.exists():
    cd = docx.Document(str(CHECKLIST))
    ms_all = " ".join(ms_words)
    pdf_all = " ".join(pdf_words)
    leaks = 0
    for p in cd.paragraphs:
        t = " ".join(words(p.text))
        if len(t) > 40 and t in pdf_all and t not in ms_all:
            leaks += 1
            say("  LEAK: " + p.text[:100])
    say(f"checklist paragraphs found in book but not in manuscript: {leaks}")
else:
    say("  (checklist file not given)")

# ---------------------------------------------------------------- page types
def folio_bottom(page):
    for b in page.get_text("dict")["blocks"]:
        if b["type"] == 0:
            for l in b["lines"]:
                s = l["spans"][0]
                if s["origin"][1] > 600 and s["text"].strip():
                    return s["text"].strip()
    return None


def head(page):
    spans = []
    for b in page.get_text("dict")["blocks"]:
        if b["type"] == 0:
            for l in b["lines"]:
                for s in l["spans"]:
                    if s["origin"][1] < 60 and s["text"].strip():
                        spans.append(s)
    return spans


say("\n== 3. Openings and TOC ==")
first_arabic = None
openers = {}
for i in range(N):
    f = folio_bottom(doc[i])
    if f and f.isdigit():
        if first_arabic is None and f == "1":
            first_arabic = i
        lines = body_lines(doc[i])
        title = [t for y, t, sp in lines if "Display-Sans" in sp[0]["font"] and
                 (sp[0]["size"] >= 20 or ("Ultra-Bold" in sp[0]["font"] and sp[0]["size"] >= 15.9))]
        openers[" ".join(title).replace("  ", " ").strip()] = int(f)
say(f"arabic page 1 = PDF page {first_arabic + 1}")
for k, v in openers.items():
    say(f"  p.{v:>3}  {k}")

toc_text = "\n".join(doc[i].get_text() for i in toc_pages)
toc_nums = [int(x) for x in re.findall(r"\.{3,}\s*(\d+)", toc_text)]
say(f"TOC page numbers: {toc_nums}")
say(f"opening folios   : {sorted(openers.values())}")
phase_pages = []
for i in range(N):
    txt = doc[i].get_text()
    if re.match(r"\s*PHASE [IV]+\s*$", txt.split("\n")[0] if txt else ""):
        phase_pages.append(i - first_arabic + 1)
say(f"phase divider pages: {phase_pages}")
expected = sorted(list(openers.values()) + phase_pages)
say("TOC matches actual pages: " + str(sorted(toc_nums) == sorted(
    [n for n in expected if n in toc_nums] + [n for n in toc_nums if n not in expected])
    and set(toc_nums) <= set(expected)))
missing = [n for n in toc_nums if n not in expected]
if missing:
    say(f"  TOC numbers not matching an opening/phase page: {missing}")

# ---------------------------------------------------------------- heads
say("\n== 4. Running heads / folios ==")
issues = 0
for i in range(first_arabic, N):
    pg = doc[i]
    n = i - first_arabic + 1
    h = head(pg)
    txt = pg.get_text().strip()
    blank = not txt
    is_opener = folio_bottom(pg) is not None
    is_phase = txt.startswith("PHASE")
    if blank or is_opener or is_phase:
        if h:
            issues += 1
            say(f"  p.{n}: unexpected running head")
        continue
    if not h:
        issues += 1
        say(f"  p.{n}: missing running head")
        continue
    folio = [s["text"].strip() for s in h if "Bold" in s["font"]]
    if not folio or str(n) not in folio[0]:
        issues += 1
        say(f"  p.{n}: folio mismatch {folio}")
    left = h[0]["bbox"][0] < 100
    if (n % 2 == 0) != left:
        issues += 1
        say(f"  p.{n}: head on wrong side")
say(f"running-head issues: {issues}")
blanks = [i - first_arabic + 1 for i in range(first_arabic, N) if not doc[i].get_text().strip()]
say(f"blank pages (arabic numbering): {blanks}")

# ---------------------------------------------------------------- stranded
say("\n== 5. Stranded headings / page-foot checks ==")
strand = 0
for i in range(N):
    lines = body_lines(doc[i])
    if not lines:
        continue
    y, t, spans = lines[-1]
    f = spans[0]["font"]
    if doc[i].get_text().strip().startswith("PHASE"):
        continue
    if ("Display-Sans" in f and spans[0]["size"] < 20) or (
            all("Bold" in s["font"] for s in spans if s["text"].strip()) and len(lines) > 3
            and y < 560):
        if i - (first_arabic or 0) + 1 > 0:
            strand += 1
            say(f"  p.{i - first_arabic + 1}: page ends with heading-like line: {t[:60]}")
say(f"stranded headings: {strand}")

# ---------------------------------------------------------------- fonts/size
say("\n== 6. Fonts and trim ==")
fonts = set()
for p in doc:
    for f in p.get_fonts():
        fonts.add((f[3], f[1]))
say("fonts: " + ", ".join(sorted(f"{n} [{t}]" for n, t in fonts)))
sizes = {(round(p.rect.width), round(p.rect.height)) for p in doc}
say(f"page sizes (pt): {sizes}   pages: {N}")

(HERE / "qc" / "qc_report.txt").write_text("\n".join(report) + "\n", encoding="utf-8")
